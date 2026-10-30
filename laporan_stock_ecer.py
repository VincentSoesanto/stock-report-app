"""Laporan Stock Ecer - desktop app to build, save (xlsx) and print the stock report."""
import hashlib
import hmac
import json
import math
import os
import re
import sqlite3
import sys
import tkinter as tk
from contextlib import closing
from datetime import datetime, timedelta
from tkinter import filedialog, messagebox, simpledialog, ttk

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.protection import hash_password as excel_hash
from openpyxl.workbook.protection import WorkbookProtection

APP_TITLE = "Laporan Stock Ecer"
DEFAULT_DB_NAME = "master_barang_contoh.xlsx"  # sample item master; pick your own with the folder button
DB_SHEET = "ACCDBRG"
MIN_PRINT_ROWS = 15
ROWS_PER_PAGE = 20  # at most this many item rows on one printed paper
MAX_DROPDOWN_ROWS = 500
MAX_COL_WIDTH = 60  # longer text wraps onto extra lines instead of widening the column
LINE_HEIGHT = 15    # points per line of Arial 11

# (key, header text, width in app px, width in xlsx chars, source field or None)
COLUMNS = [
    ("kode", "Kode Barang", 110, 18, "kode"),
    ("nama", "Nama Barang", 230, 32, "nama"),
    ("stn", "StnPerPak", 80, 11, "stn"),
    ("satuan", "Satuan", 70, 10, "satuan"),
    ("koli", "J. Koli\nPermintaan", 125, 12, None),
    ("gudang", "Gudang", 80, 12, None),
    ("ttd_gudang", "TTD Kpl.\nGudang", 110, 12, None),
    ("opname", "Stock\nOpname", 105, 12, None),
    ("ttd_ecer", "TTD Ecer", 80, 12, None),
    ("ttd_kantor", "TTD Kantor", 85, 12, None),
]
# hand columns filled by picking from a list instead of typing ("" = empty, to clear the cell)
CHOICES = {
    "gudang": ["", "GD1", "GD2", "GD3"],
}
DB_KEYS = [c[0] for c in COLUMNS if c[4]]
HAND_KEYS = [c[0] for c in COLUMNS if not c[4]]  # typed in by the user (or left blank and written on paper)
EDIT_KEYS = [k for k in HAND_KEYS if not k.startswith("ttd_")]  # TTD columns are only signed on paper


def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


CONFIG_PATH = os.path.join(app_dir(), "laporan_config.json")
DB_PATH = os.path.join(app_dir(), "laporan.db")
# old approver list ('password = nama' per line); read once into laporan_config.json, then no longer used
APPROVERS_PATH = os.path.join(app_dir(), "password_disetujui.txt")
# settings password (gear window, Impor Laporan Lama, unlocking an approved printout): this one until it is
# changed in the gear window, then only its hash is kept in laporan_config.json ("settings_password_hash")
DEFAULT_SETTINGS_PASSWORD = "admin123"  # demo default: change it in the gear window after the first start
ROW_COLORS = ("white", "#e6e6e6")  # main table: odd rows white, even rows light gray
REPORT_NAME_RE = re.compile(r"^Laporan Stock Ecer (\d{8}_\d{6})\.xlsx$", re.IGNORECASE)


def load_config():
    try:
        with open(CONFIG_PATH, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_config(cfg):
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2)
    except OSError:
        pass


def cell_text(v):
    """Value as shown in a table: None -> '', 12.0 -> '12'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def cell_value(text):
    """Typed text for the xlsx: numbers become real numbers so Excel can add them up."""
    s = str(text).strip()
    if re.fullmatch(r"-?\d+", s) and not (len(s) > 1 and s.startswith("0")):
        return int(s)
    if re.fullmatch(r"-?\d+[.,]\d+", s):
        return float(s.replace(",", "."))
    return s


def read_old_approvers():
    """Entries of the old password_disetujui.txt as [{'nama', 'password'}], or [] if there is no file."""
    approvers, seen = [], set()
    try:
        with open(APPROVERS_PATH, encoding="utf-8-sig") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                pw, name = (s.strip() for s in line.split("=", 1))
                if pw and name and pw not in seen:
                    seen.add(pw)
                    approvers.append({"nama": name, "password": pw})
    except OSError:
        pass
    return approvers


PBKDF2_ROUNDS = 100_000


def hash_secret(pw):
    """One-way hash of a password with a random salt: 'pbkdf2_sha256$rounds$salt$hash'."""
    salt = os.urandom(16)
    digest = hashlib.pbkdf2_hmac("sha256", pw.encode("utf-8"), salt, PBKDF2_ROUNDS)
    return f"pbkdf2_sha256${PBKDF2_ROUNDS}${salt.hex()}${digest.hex()}"


def check_secret(pw, stored):
    """True if pw is the password behind a hash_secret() value."""
    try:
        _name, rounds, salt, digest = stored.split("$")
        calc = hashlib.pbkdf2_hmac("sha256", pw.encode("utf-8"), bytes.fromhex(salt), int(rounds))
    except (AttributeError, ValueError):
        return False
    return hmac.compare_digest(calc.hex(), digest)


def get_approvers(cfg):
    """Approver list [{'nama', 'hash'}] kept in laporan_config.json. The first time, it is filled from the old
    txt file; plain passwords (older versions) are replaced by their hash."""
    changed = "approvers" not in cfg
    if changed:
        cfg["approvers"] = read_old_approvers()
    for a in cfg["approvers"]:
        if "password" in a:
            a["hash"] = hash_secret(a.pop("password"))
            changed = True
    if changed:
        save_config(cfg)
    return cfg["approvers"]


def check_settings_password(cfg, pw):
    """True if pw is the settings password (the changed one if set, else the default)."""
    stored = cfg.get("settings_password_hash")
    return check_secret(pw, stored) if stored else pw == DEFAULT_SETTINGS_PASSWORD


def find_approver(cfg, pw):
    """Name of the approver whose password is pw, or ''."""
    return next((a["nama"] for a in get_approvers(cfg) if pw and check_secret(pw, a["hash"])), "")


def upgrade_config(cfg):
    """Turn passwords that older versions saved as plain text into hashes."""
    get_approvers(cfg)
    if "password" in cfg:  # Excel edit lock: keep only the hash Excel itself stores in the file
        pw = cfg.pop("password")
        if pw:
            cfg["excel_password_hash"] = excel_hash(pw)
        save_config(cfg)


def text_width(text):
    """Approximate width of text in Excel column units for Arial 11 (capitals and digits are wider)."""
    w = 0.0
    for ch in text:
        if ch.isupper() or ch.isdigit() or ch in "MW@%&":
            w += 1.4
        elif ch in " .,:;'!|il()[]/-":
            w += 0.7
        else:
            w += 1.05
    return w


def row_height(texts, widths, minimum):
    """Row height in points so every (wrapped) text in the row is fully visible."""
    lines = 1
    for text, width in zip(texts, widths):
        usable = max(width - 1, 1)
        n = sum(max(1, math.ceil(text_width(part) / usable)) for part in str(text).split("\n"))
        lines = max(lines, n)
    return max(minimum, lines * LINE_HEIGHT + 5)


# ---------- watermark ----------
# Excel's own watermark: a picture in the page header (&G). It is drawn behind the cells on every printed page,
# but Excel shows header pictures only in print preview / Page Layout. So the same picture is also the sheet
# background, which shows in the normal view (backgrounds are never printed, so it isn't on paper twice).
# openpyxl can't add either, so add_watermark() puts them into the saved xlsx file afterwards.
WATERMARK_TEXT = "NamaToko"  # default shop name; set "watermark_text" in laporan_config.json to change it
PAGE_W_PT, PAGE_H_PT = 842, 595  # A4 landscape, the whole sheet of paper
HEADER_MARGIN_IN = 0.1  # the printed watermark starts this far from the top edge of the paper


# very light gray = low opacity, black text stays easy to read; Excel shows the screen background darker
# than the picture itself, so that one is lighter
WATERMARK_GRAY_PRINT = (227, 227, 227)
WATERMARK_GRAY_SCREEN = (241, 241, 241)


def watermark_text():
    """Shop name for the watermark: "watermark_text" in laporan_config.json, else WATERMARK_TEXT."""
    return str(load_config().get("watermark_text") or WATERMARK_TEXT).strip() or WATERMARK_TEXT


def watermark_image(w, h, unit=1.0, gray=WATERMARK_GRAY_PRINT):
    """Picture w x h px: the watermark text many times over, small, slanted, in rows that are shifted
    against each other (a pattern is harder to fake than one big word). unit = px per screen pixel, so
    the text is the same size on screen and on paper."""
    from PIL import Image, ImageDraw, ImageFont
    size = round(30 * unit)
    try:
        font = ImageFont.truetype("arialbd.ttf", size)
    except OSError:
        font = ImageFont.load_default(size=size)
    word = watermark_text()
    box = font.getbbox(word)
    text = Image.new("L", (box[2] + 8, box[3] + 8), 0)
    ImageDraw.Draw(text).text((4, 4), word, font=font, fill=255)
    text = text.rotate(25, expand=True, resample=Image.BICUBIC)
    img = Image.new("RGB", (w, h), "white")
    ink = Image.new("RGB", text.size, gray)
    step_x, step_y = text.width + round(40 * unit), round(110 * unit)
    for row, y in enumerate(range(-text.height // 2, h + step_y, step_y)):  # one row/column extra: no gaps
        for x in range(-text.width + (row % 2) * step_x // 2, w + step_x, step_x):  # at the edges
            img.paste(ink, (x, y), text)
    return img


def png_bytes(img):
    import io
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


_VML_HEADER = """<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" \
xmlns:x="urn:schemas-microsoft-com:office:excel">
 <o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="1"/></o:shapelayout>
 <v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75" o:preferrelative="t"
  path="m@4@5l@4@11@9@11@9@5xe" filled="f" stroked="f">
  <v:stroke joinstyle="miter"/>
  <v:formulas>
   <v:f eqn="if lineDrawn pixelLineWidth 0"/><v:f eqn="sum @0 1 0"/><v:f eqn="sum 0 0 @1"/>
   <v:f eqn="prod @2 1 2"/><v:f eqn="prod @3 21600 pixelWidth"/><v:f eqn="prod @3 21600 pixelHeight"/>
   <v:f eqn="sum @0 0 1"/><v:f eqn="prod @6 1 2"/><v:f eqn="prod @7 21600 pixelWidth"/>
   <v:f eqn="sum @8 21600 0"/><v:f eqn="prod @7 21600 pixelHeight"/><v:f eqn="sum @10 21600 0"/>
  </v:formulas>
  <v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/>
  <o:lock v:ext="edit" aspectratio="t"/>
 </v:shapetype>
 <v:shape id="CH" o:spid="_x0000_s1025" type="#_x0000_t75"
  style="position:absolute;margin-left:0;margin-top:0;width:{w}pt;height:{h}pt;z-index:1">
  <v:imagedata o:relid="rId1" o:title="watermark"/>
  <o:lock v:ext="edit" rotation="t"/>
 </v:shape>
</xml>"""
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_RELS = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' \
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' \
        '<Relationship Id="rId1" Type="' + _REL + '/{type}" Target="{target}"/>{more}</Relationships>'


def add_watermark(path, table_px):
    """Put the watermark into the first sheet of a saved xlsx made by build_workbook (header '&C&G'):
    as sheet background on screen, the size of the table (title row to last row, columns A-J;
    table_px = (width, height) in screen pixels), and as header picture on paper, filling the whole
    printable page. The file is rewritten in place."""
    import zipfile
    from PIL import Image
    tw, th = table_px
    # screen: the background repeats over the whole sheet, so the picture is the table with a wide white
    # border right and below; the watermark itself only covers the table
    bg = Image.new("RGB", (max(tw + 2000, 4000), max(th + 2000, 4000)), "white")
    bg.paste(watermark_image(tw, th, gray=WATERMARK_GRAY_SCREEN), (0, 0))
    # paper: the whole sheet of paper edge to edge, every page (the header picture repeats; it starts
    # HEADER_MARGIN_IN from the top and is centred). A fixed size in points that doesn't scale with the page
    # (scaleWithDoc off in build_workbook): sizes from screen pixels came out too small on another laptop,
    # and only filling inside the margins left the edges empty. The printer's own unprintable edge stays white.
    pic_w, pic_h = PAGE_W_PT, round(PAGE_H_PT - HEADER_MARGIN_IN * 72)
    header = watermark_image(pic_w * 2, pic_h * 2, unit=1.5)  # 2 px per pt, 0.75 pt per screen pixel
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    sheet = "xl/worksheets/sheet1.xml"
    if "xl/worksheets/_rels/sheet1.xml.rels" in parts:  # not made by this app; leave it alone
        return
    xml = parts[sheet].decode("utf-8")
    # rId1 = header picture (printed), rId2 = sheet background (on screen); this order is required by Excel
    xml = xml.replace("</worksheet>", f'<legacyDrawingHF xmlns:r="{_REL}" r:id="rId1"/>'
                                      f'<picture xmlns:r="{_REL}" r:id="rId2"/></worksheet>')
    parts[sheet] = xml.encode("utf-8")
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = _RELS.format(
        type="vmlDrawing", target="../drawings/vmlDrawingHF1.vml",
        more=f'<Relationship Id="rId2" Type="{_REL}/image" Target="../media/watermark_bg.png"/>').encode("utf-8")
    parts["xl/drawings/vmlDrawingHF1.vml"] = _VML_HEADER.replace("{w}", str(pic_w)).replace(
        "{h}", str(pic_h)).encode("utf-8")
    parts["xl/drawings/_rels/vmlDrawingHF1.vml.rels"] = _RELS.format(
        type="image", target="../media/watermark.png", more="").encode("utf-8")
    parts["xl/media/watermark.png"] = png_bytes(header)
    parts["xl/media/watermark_bg.png"] = png_bytes(bg)
    types = parts["[Content_Types].xml"].decode("utf-8")
    for ext, ctype in (("png", "image/png"), ("vml", "application/vnd.openxmlformats-officedocument.vmlDrawing")):
        if f'Extension="{ext}"' not in types:
            types = types.replace("<Default ", f'<Default Extension="{ext}" ContentType="{ctype}"/><Default ', 1)
    parts["[Content_Types].xml"] = types.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def save_report(wb, path):
    """Save a report workbook from build_workbook, then add the watermark the size of its table."""
    wb.save(path)
    add_watermark(path, wb.table_px)


def load_items(path):
    """Read items from the database workbook: B=kode, C=nama, D=satuan, F=stnperpak."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[DB_SHEET] if DB_SHEET in wb.sheetnames else wb.worksheets[0]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * 6
        kode = row[1]
        if kode is None or str(kode).strip() == "":
            continue
        stn = row[5]
        if isinstance(stn, float) and stn.is_integer():
            stn = int(stn)
        item = {
            "kode": str(kode).strip(),
            "nama": "" if row[2] is None else str(row[2]).strip(),
            "satuan": "" if row[3] is None else str(row[3]).strip(),
            "stn": "" if stn is None else stn,
        }
        item["_search"] = " ".join(str(item[k]) for k in DB_KEYS).lower()
        items.append(item)
    wb.close()
    return items


# ---------- history database (SQLite, one file next to the exe) ----------
def db_connect():
    con = sqlite3.connect(DB_PATH)
    con.execute("PRAGMA foreign_keys = ON")
    con.executescript("""
        CREATE TABLE IF NOT EXISTS reports (
            id INTEGER PRIMARY KEY,
            printed_at TEXT NOT NULL,          -- YYYY-MM-DD HH:MM:SS
            perwakilan_ecer TEXT,
            perwakilan_kantor TEXT,
            file_name TEXT UNIQUE,             -- also stops importing the same file twice
            file_path TEXT
        );
        CREATE TABLE IF NOT EXISTS report_items (
            report_id INTEGER NOT NULL REFERENCES reports(id) ON DELETE CASCADE,
            line_no INTEGER NOT NULL,
            kode TEXT NOT NULL,
            nama TEXT,
            satuan TEXT,
            stn TEXT,
            PRIMARY KEY (report_id, line_no)
        );
        CREATE INDEX IF NOT EXISTS idx_reports_time ON reports(printed_at);
        CREATE INDEX IF NOT EXISTS idx_items_kode ON report_items(kode);
    """)
    # hand-filled columns (J. Koli Permintaan ... TTD Kantor); added to databases made before they existed
    have = {r[1] for r in con.execute("PRAGMA table_info(reports)")}
    if "disetujui" not in have:
        con.execute("ALTER TABLE reports ADD COLUMN disetujui TEXT")
    if "file_mtime" not in have:  # 'last changed' time of the xlsx when Riwayat last matched it
        con.execute("ALTER TABLE reports ADD COLUMN file_mtime REAL")
    have = {r[1] for r in con.execute("PRAGMA table_info(report_items)")}
    for k in HAND_KEYS:
        if k not in have:
            con.execute(f"ALTER TABLE report_items ADD COLUMN {k} TEXT")
    return con


def db_save_report(printed_at, ecer, kantor, path, rows, disetujui=""):
    """Record one printed report and its items. Returns False if that file is already recorded."""
    with closing(db_connect()) as con, con:
        cur = con.execute(
            "INSERT OR IGNORE INTO reports (printed_at, perwakilan_ecer, perwakilan_kantor, disetujui, file_name,"
            " file_path, file_mtime) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (printed_at.strftime("%Y-%m-%d %H:%M:%S"), ecer, kantor, disetujui, os.path.basename(path), path,
             file_mtime(path)))
        if cur.rowcount == 0:
            return False
        keys = ["kode", "nama", "satuan", "stn"] + HAND_KEYS
        con.executemany(
            f"INSERT INTO report_items (report_id, line_no, {', '.join(keys)})"
            f" VALUES (?, ?, {', '.join('?' * len(keys))})",
            [(cur.lastrowid, i, *(str(r.get(k, "")) for k in keys)) for i, r in enumerate(rows, 1)])
    return True


def file_mtime(path):
    """'Last changed' time of a file, or None if it is missing."""
    try:
        return os.path.getmtime(path)
    except OSError:
        return None


def db_update_items(report_id, changes, mtime):
    """Save the values changed in Riwayat ({line_no: {key: value}}) in one go, with the rewritten file's time."""
    with closing(db_connect()) as con, con:
        for line_no, values in changes.items():
            for key, value in values.items():
                if key not in HAND_KEYS:
                    raise ValueError(key)
                con.execute(f"UPDATE report_items SET {key} = ? WHERE report_id = ? AND line_no = ?",
                            (value, report_id, line_no))
        con.execute("UPDATE reports SET file_mtime = ? WHERE id = ?", (mtime, report_id))


def db_replace_report(con, report_id, ecer, kantor, disetujui, rows, path):
    """Overwrite a printout in Riwayat with what its xlsx file says now."""
    con.execute("UPDATE reports SET perwakilan_ecer = ?, perwakilan_kantor = ?, disetujui = ?, file_path = ?,"
                " file_mtime = ? WHERE id = ?", (ecer, kantor, disetujui, path, file_mtime(path), report_id))
    con.execute("DELETE FROM report_items WHERE report_id = ?", (report_id,))
    keys = ["kode", "nama", "satuan", "stn"] + HAND_KEYS
    con.executemany(
        f"INSERT INTO report_items (report_id, line_no, {', '.join(keys)}) VALUES (?, ?, {', '.join('?' * len(keys))})",
        [(report_id, i, *(str(r.get(k, "")) for k in keys)) for i, r in enumerate(rows, 1)])


def db_report_content(con, report_id):
    """(ecer, kantor, disetujui, rows) of a printout as stored, in the same shape read_report_file gives."""
    ecer, kantor, disetujui = con.execute("SELECT perwakilan_ecer, perwakilan_kantor, disetujui FROM reports"
                                          " WHERE id = ?", (report_id,)).fetchone()
    keys = [c[0] for c in COLUMNS]
    rows = [{k: cell_text(v) for k, v in zip(keys, r)} for r in con.execute(
        f"SELECT {', '.join(keys)} FROM report_items WHERE report_id = ? ORDER BY line_no", (report_id,))]
    return ecer or "", kantor or "", disetujui or "", rows


def sync_report_file(con, report_id, path, stored_mtime, printed_at, force=False):
    """Bring a printout in Riwayat up to date with its xlsx file when the file was changed outside the app
    (edited in Excel). Returns True if Riwayat was updated. force = check the file even if its time is the same."""
    mtime = file_mtime(path)
    if mtime is None or (stored_mtime is not None and abs(mtime - stored_mtime) < 0.01 and not force):
        return False
    if stored_mtime is None and not force:
        # recorded by an older version that did not keep the time: the file was edited only if it changed
        # well after printing; otherwise Riwayat (maybe corrected there) is kept and only the time is noted
        printed = datetime.strptime(printed_at, "%Y-%m-%d %H:%M:%S").timestamp()
        if mtime < printed + 120:
            con.execute("UPDATE reports SET file_mtime = ? WHERE id = ?", (mtime, report_id))
            return False
    try:
        ecer, kantor, rows, disetujui = read_report_file(path)
    except Exception:  # noqa: BLE001 - open in Excel right now, damaged, ...: try again next time
        return False
    if (ecer, kantor, disetujui, rows) == db_report_content(con, report_id):
        con.execute("UPDATE reports SET file_mtime = ?, file_path = ? WHERE id = ?", (mtime, path, report_id))
        return False
    db_replace_report(con, report_id, ecer, kantor, disetujui, rows, path)
    return True


def db_sync_files(date_from, date_to):
    """Check every printout between the dates against its xlsx file. Returns how many were updated."""
    params = (date_from.strftime("%Y-%m-%d"), (date_to + timedelta(days=1)).strftime("%Y-%m-%d"))
    updated = 0
    with closing(db_connect()) as con, con:
        for report_id, path, stored, printed_at in con.execute(
                "SELECT id, file_path, file_mtime, printed_at FROM reports"
                " WHERE printed_at >= ? AND printed_at < ?", params).fetchall():
            updated += sync_report_file(con, report_id, path, stored, printed_at)
    return updated


def read_report_file(path):
    """Read Perwakilan Ecer/Kantor and the item rows back from a report xlsx made by this app."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        values = [list(r) + [None] * len(COLUMNS) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()

    def label(row):  # column B holds ": value"
        v = values[row - 1][1] if len(values) >= row else None
        return "" if v is None else str(v).lstrip(":").strip()

    rows = []
    for r in values[7:]:  # table data starts at row 8
        if r[0] is None or str(r[0]).strip() == "":
            continue
        rows.append({c[0]: cell_text(r[i]) for i, c in enumerate(COLUMNS)})
    return label(4), label(5), rows, label(6)


def import_report_files(paths):
    """Add report files ('Laporan Stock Ecer YYYYMMDD_HHMMSS.xlsx') to the history.
    A report already in the history is updated from the file when the file is different (edited in Excel).
    Returns (print times added, number updated, number already the same, names that are not a report,
    failed names)."""
    added, updated, same, others, failed = [], 0, 0, [], []
    with closing(db_connect()) as con:
        known = {r[0]: r[1:] for r in con.execute("SELECT file_name, id, file_mtime, printed_at FROM reports")}
    for path in sorted(paths):
        name = os.path.basename(path)
        m = REPORT_NAME_RE.match(name)
        if not m:  # another xlsx
            others.append(name)
            continue
        if name in known:
            report_id, stored, printed_at = known[name]
            with closing(db_connect()) as con, con:
                if sync_report_file(con, report_id, path, stored, printed_at, force=True):
                    updated += 1
                else:
                    same += 1
            continue
        try:
            ecer, kantor, rows, disetujui = read_report_file(path)
            printed_at = datetime.strptime(m.group(1), "%Y%m%d_%H%M%S")
            if db_save_report(printed_at, ecer, kantor, path, rows, disetujui):
                added.append(printed_at)
        except Exception:  # noqa: BLE001 - skip unreadable files, report them to the user
            failed.append(name)
    return added, updated, same, others, failed


def import_report_folder(folder):
    """Add every report file in folder to the history (first start). Returns (added, failed)."""
    if not os.path.isdir(folder):
        return 0, []
    added, _updated, _same, _others, failed = import_report_files(
        [os.path.join(folder, n) for n in os.listdir(folder) if n.lower().endswith(".xlsx")])
    return len(added), failed


def _history_filter(date_from, date_to, terms):
    """WHERE parts + params: printed between date_from and date_to (inclusive), item kode/nama has every word."""
    where = ["r.printed_at >= ?", "r.printed_at < ?"]
    params = [date_from.strftime("%Y-%m-%d"), (date_to + timedelta(days=1)).strftime("%Y-%m-%d")]
    for t in terms:
        t = t.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        where.append("(i.kode || ' ' || COALESCE(i.nama, '')) LIKE ? ESCAPE '\\'")
        params.append(f"%{t}%")
    return where, params


def db_prints(date_from, date_to, terms):
    """One row per printout that has a matching item: printed_at, perwakilan ecer, perwakilan kantor,
    disetujui, number of items, then hidden: file_path, report_id."""
    where, params = _history_filter(date_from, date_to, terms)
    sql = ("SELECT r.printed_at, r.perwakilan_ecer, r.perwakilan_kantor, r.disetujui,"
           " (SELECT COUNT(*) FROM report_items n WHERE n.report_id = r.id), r.file_path, r.id"
           " FROM reports r WHERE r.printed_at >= ? AND r.printed_at < ?")
    if terms:  # the date check is also inside, so the same params fit
        sql += (" AND EXISTS (SELECT 1 FROM report_items i WHERE i.report_id = r.id AND "
                + " AND ".join(where) + ")")
        params = params[:2] + params
    with closing(db_connect()) as con:
        return con.execute(sql + " ORDER BY r.printed_at DESC", params).fetchall()


def db_item_totals(date_from, date_to, terms):
    """One row per item: kode, nama, stn, satuan, total J. Koli Permintaan (numbers only), times printed."""
    where, params = _history_filter(date_from, date_to, terms)
    sql = ("SELECT i.kode, i.nama, i.stn, i.satuan,"
           " SUM(CASE WHEN TRIM(i.koli) GLOB '[0-9]*' THEN CAST(TRIM(i.koli) AS REAL) END), COUNT(*)"
           " FROM report_items i JOIN reports r ON r.id = i.report_id WHERE " + " AND ".join(where) +
           " GROUP BY i.kode ORDER BY COUNT(*) DESC, i.kode")
    with closing(db_connect()) as con:
        return con.execute(sql, params).fetchall()


def db_print_row(report_id):
    """One printout in the same shape as a db_prints row."""
    with closing(db_connect()) as con:
        return con.execute("SELECT r.printed_at, r.perwakilan_ecer, r.perwakilan_kantor, r.disetujui,"
                           " (SELECT COUNT(*) FROM report_items n WHERE n.report_id = r.id), r.file_path, r.id"
                           " FROM reports r WHERE r.id = ?", (report_id,)).fetchone()


def db_sync_one(report_id):
    """Check one printout against its xlsx file. True if Riwayat was updated."""
    with closing(db_connect()) as con, con:
        path, stored, printed_at = con.execute("SELECT file_path, file_mtime, printed_at FROM reports WHERE id = ?",
                                               (report_id,)).fetchone()
        return sync_report_file(con, report_id, path, stored, printed_at)


def db_report_items(report_id):
    """Every line of one printout: line_no, then the 10 report columns."""
    with closing(db_connect()) as con:
        return con.execute(f"SELECT line_no, {', '.join(c[0] for c in COLUMNS)} FROM report_items"
                           " WHERE report_id = ? ORDER BY line_no", (report_id,)).fetchall()


def fmt_time(s):
    """'2026-09-30 11:19:05' -> '30/09/2026 11:19'"""
    return datetime.strptime(s, "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y %H:%M")


def sort_key(v):
    """Numbers (also numeric text like '100') before text, text case-insensitive."""
    if v is None:
        return (1, "")
    try:
        return (0, float(v))
    except (TypeError, ValueError):
        return (1, str(v).lower())


def db_start_new():
    """Move the current history to a dated backup file so a new, empty history starts. Returns the backup path."""
    backup = os.path.join(os.path.dirname(DB_PATH), f"laporan_backup_{datetime.now():%Y%m%d_%H%M%S}.db")
    os.replace(DB_PATH, backup)
    db_connect().close()  # create the new empty database
    return backup


class HistoryWindow(tk.Toplevel):
    """Riwayat: a list of every printout (or totals per item), searchable by item and date range.
    Double-clicking a printout shows everything on it; Kembali goes back to the list."""

    PRINTS = [("waktu", "Tanggal / Jam Print", 140, "center"), ("p_ecer", "Perwakilan Ecer", 170, "w"),
              ("p_kantor", "Perwakilan Kantor", 170, "w"), ("disetujui", "Disetujui", 150, "w"),
              ("jumlah", "Jumlah Barang", 110, "center")]
    SUMMARY = [("kode", "Kode Barang", 110, "w"), ("nama", "Nama Barang", 300, "w"),
               ("stn", "StnPerPak", 80, "center"), ("satuan", "Satuan", 70, "center"),
               ("koli", "Total J. Koli Permintaan", 150, "center"), ("jumlah", "Jumlah Dicetak", 110, "center")]
    DESC_FIRST = {"summary": {4, 5}, "prints": {0, 4}}  # columns that sort biggest/newest first
    DAY_COLORS = ("white", "#e6e6e6")  # printout rows switch colour whenever the day changes

    def __init__(self, master):
        super().__init__(master)
        self.app = master
        self.title(f"Riwayat - {APP_TITLE}")
        self.geometry("1280x600")
        self.minsize(800, 380)
        self.transient(master)
        self._after = None
        self._rows = []
        self._report = None  # the printout row shown on the view page, None while the list shows
        self._pending = {}  # changes typed on the view page, not saved yet: {line_no: {key: value}}
        # (column index, descending) per mode; default = newest print first / most printed first
        self._sort = {"summary": (5, True), "prints": (0, True)}
        self._build_list_page()
        self._build_view_page()
        self.list_page.pack(fill="both", expand=True)

        self.bind("<Escape>", lambda e: self.back() if self._report else self.close())
        self.protocol("WM_DELETE_WINDOW", self.close)
        self.set_month(0)
        self.refresh(sync=True)  # once on opening: pick up report files edited in Excel
        self.search.focus_set()

    # ---------- list page ----------
    def _build_list_page(self):
        page = self.list_page = tk.Frame(self)
        top = tk.Frame(page)
        top.pack(fill="x", padx=12, pady=(10, 4))
        tk.Label(top, text="Cari barang:").grid(row=0, column=0, sticky="w")
        self.search_var = tk.StringVar()
        self.search = ttk.Entry(top, textvariable=self.search_var, width=34)
        self.search.grid(row=0, column=1, sticky="w", padx=(4, 16))
        tk.Label(top, text="Dari").grid(row=0, column=2)
        self.from_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.from_var, width=11).grid(row=0, column=3, padx=4)
        tk.Label(top, text="Sampai").grid(row=0, column=4)
        self.to_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.to_var, width=11).grid(row=0, column=5, padx=4)
        ttk.Button(top, text="Bulan Ini", command=lambda: self.set_month(0)).grid(row=0, column=6, padx=(8, 2))
        ttk.Button(top, text="Bulan Lalu", command=lambda: self.set_month(-1)).grid(row=0, column=7, padx=2)
        self._refresh_img = refresh_icon(self)  # keep a reference for Tk
        b = ttk.Button(top, image=self._refresh_img, command=lambda: self.refresh(sync=True))
        b.grid(row=0, column=8, padx=(12, 2))
        Tooltip(b, "Muat ulang: cek lagi file Excel laporan yang diubah di Excel")

        self.mode_var = tk.StringVar(value="prints")
        modes = tk.Frame(page)
        modes.pack(fill="x", padx=12)
        ttk.Radiobutton(modes, text="Daftar print (satu baris per print)", value="prints",
                        variable=self.mode_var, command=self.refresh).pack(side="left")
        ttk.Radiobutton(modes, text="Total per barang (dijumlah)", value="summary",
                        variable=self.mode_var, command=self.refresh).pack(side="left", padx=16)
        self.hint = tk.Label(modes, text="", fg="#555")
        self.hint.pack(side="left", padx=8)

        self.tree = self._table(page)
        self.tree.bind("<Double-1>", self._on_double_click)
        self.tree.bind("<Return>", lambda e: self._open_selected())
        for i, color in enumerate(self.DAY_COLORS):
            self.tree.tag_configure(f"day{i}", background=color)

        bottom = tk.Frame(page)
        bottom.pack(fill="x", padx=12, pady=(0, 10))
        self.info = tk.Label(bottom, text="", anchor="w")
        self.info.pack(side="left", fill="x", expand=True)
        ttk.Button(bottom, text="Tutup", command=self.close).pack(side="right", padx=3)
        ttk.Button(bottom, text="Impor Laporan Lama...", command=self.import_old).pack(side="right", padx=3)
        ttk.Button(bottom, text="Mulai Riwayat Baru...", command=self.start_new).pack(side="right", padx=3)
        print_button(bottom, self.print_selected).pack(side="right", padx=(3, 12))

        for var in (self.search_var, self.from_var, self.to_var):
            var.trace_add("write", lambda *a: self._schedule())

    def _table(self, page):
        """Report-style table with scrollbars, packed into page."""
        frame = tk.Frame(page)
        frame.pack(fill="both", expand=True, padx=12, pady=6)
        tree = ttk.Treeview(frame, show="headings", style="Report.Treeview", selectmode="browse")
        ys = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        xs = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ys.grid(row=0, column=1, sticky="ns")
        xs.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        tree.bind("<MouseWheel>", lambda e: close_editors(tree), add="+")
        return tree

    def close(self):
        close_editors(self.vtree)  # keep a value still being typed
        if self._confirm_pending():
            self.destroy()

    def set_month(self, offset):
        first = datetime.now().replace(day=1)
        last = datetime.now()
        if offset == -1:
            last = first - timedelta(days=1)
            first = last.replace(day=1)
        self.from_var.set(first.strftime("%d/%m/%Y"))
        self.to_var.set(last.strftime("%d/%m/%Y"))
        self.refresh()

    def _schedule(self):
        if self._after:
            self.after_cancel(self._after)
        self._after = self.after(250, self.refresh)

    def refresh(self, sync=False):
        """Fill the list. sync = first re-read report files changed in Excel (on opening and with the
        refresh button; not on every search letter, to stay fast)."""
        self._after = None
        try:
            d_from = datetime.strptime(self.from_var.get().strip(), "%d/%m/%Y")
            d_to = datetime.strptime(self.to_var.get().strip(), "%d/%m/%Y")
        except ValueError:
            self.info.config(text="Tanggal tidak valid, gunakan format dd/mm/yyyy.", fg="#c00")
            return
        summary = self.mode_var.get() == "summary"
        terms = self.search_var.get().lower().split()
        try:
            synced = db_sync_files(d_from, d_to) if sync else 0  # files edited in Excel -> Riwayat
            rows = (db_item_totals if summary else db_prints)(d_from, d_to, terms)
        except sqlite3.Error as e:
            self.info.config(text=f"Gagal membaca riwayat: {e}", fg="#c00")
            return

        self._rows = [list(r) for r in rows]
        self._render()
        if summary:
            self.hint.config(text="Klik dua kali pada barang untuk melihat print yang berisi barang itu.")
            text = f"{len(rows)} barang, total {sum(r[5] for r in rows)} kali dicetak"
        else:
            self.hint.config(text="Klik dua kali pada print untuk melihat isinya.")
            text = f"{len(rows)} kali print" + (" yang berisi barang yang dicari" if terms else "")
        if synced:
            text += f"  -  {synced} laporan diperbarui dari file Excel yang diubah"
        elif sync:
            text += "  -  semua sama dengan file Excel"
        self.info.config(text=f"{text}  ({self.from_var.get()} - {self.to_var.get()})", fg="#333")

    def _render(self):
        """Fill the table from self._rows, sorted by the column whose header was clicked."""
        mode = self.mode_var.get()
        summary = mode == "summary"
        cols = self.SUMMARY if summary else self.PRINTS
        col, reverse = self._sort[mode]
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = [c[0] for c in cols]
        for i, (key, text, width, anchor) in enumerate(cols):
            arrow = ("  ▼" if reverse else "  ▲") if i == col else ""
            self.tree.heading(key, text=text + arrow, command=lambda i=i: self.sort_by(i))
            self.tree.column(key, width=width, minwidth=50, anchor=anchor)
        # column i of the table is value i of the query row, so sort on the raw value
        # (numbers as numbers, dates as YYYY-MM-DD so they sort in time order)
        order = sorted(range(len(self._rows)), key=lambda n: sort_key(self._rows[n][col]), reverse=reverse)
        band, prev_day = 0, None
        for n in order:
            r = self._rows[n]
            if summary:
                self.tree.insert("", "end", iid=str(n), values=[*r[:4], cell_text(r[4]), r[5]])
                continue
            day = r[0][:10]
            if prev_day is not None and day != prev_day:
                band = 1 - band
            prev_day = day
            self.tree.insert("", "end", iid=str(n), values=[fmt_time(r[0]), *(cell_text(v) for v in r[1:5])],
                             tags=(f"day{band}",))

    def sort_by(self, col):
        mode = self.mode_var.get()
        cur, reverse = self._sort[mode]
        # first click: counts and dates biggest/newest first, text A-Z; next click reverses
        self._sort[mode] = (col, not reverse) if col == cur else (col, col in self.DESC_FIRST[mode])
        self._render()

    def _on_double_click(self, event):
        if self.tree.identify_region(event.x, event.y) == "cell":
            self._open_selected(self.tree.identify_row(event.y))

    def _open_selected(self, iid=None):
        iid = iid or (self.tree.selection() or [None])[0]
        if not iid:
            return
        if self.mode_var.get() == "summary":
            # list the printouts that have this item
            self.search_var.set(self.tree.set(iid, "kode"))
            self.mode_var.set("prints")
            self.refresh()
            return
        self.show_print(self._rows[int(iid)])

    # ---------- view page: one printout ----------
    def _build_view_page(self):
        page = self.view_page = tk.Frame(self)
        top = tk.Frame(page)
        top.pack(fill="x", padx=12, pady=(10, 4))
        ttk.Button(top, text="◀  Kembali ke daftar print", command=self.back).pack(side="left", anchor="n")
        info = tk.Frame(top)
        info.pack(side="left", padx=(24, 0))
        self.view_vars = {}
        for i, (key, label) in enumerate([("waktu", "Tanggal / Jam Print"), ("p_ecer", "Perwakilan Ecer"),
                                          ("p_kantor", "Perwakilan Kantor"), ("disetujui", "Disetujui oleh")]):
            tk.Label(info, text=label, width=18, anchor="w").grid(row=i, column=0, sticky="w")
            self.view_vars[key] = tk.StringVar()
            tk.Label(info, textvariable=self.view_vars[key], anchor="w",
                     font=("Segoe UI", 9, "bold")).grid(row=i, column=1, sticky="w")
        self.view_hint = tk.Label(page, text="", fg="#555")
        self.view_hint.pack(anchor="w", padx=12, pady=(6, 0))

        self.vtree = self._table(page)
        keys = [c[0] for c in COLUMNS]
        self.vtree["columns"] = keys
        for key, text, width, _xw, _src in COLUMNS:
            self.vtree.heading(key, text=text.replace("\n", " "))
            self.vtree.column(key, width=width, minwidth=50, anchor="w" if key in ("kode", "nama") else "center")
        for i, color in enumerate(ROW_COLORS):
            self.vtree.tag_configure(f"row{i}", background=color)
        self.vtree.bind("<Double-1>", self._on_view_double_click)

        bottom = tk.Frame(page)
        bottom.pack(fill="x", padx=12, pady=(0, 10))
        self.view_info = tk.Label(bottom, text="", anchor="w")
        self.view_info.pack(side="left", fill="x", expand=True)
        ttk.Button(bottom, text="Tutup", command=self.close).pack(side="right", padx=3)
        ttk.Button(bottom, text="Buka File Excel", command=self.open_file).pack(side="right", padx=3)
        print_button(bottom, lambda: self.print_report(self._report)).pack(side="right", padx=3)
        tk.Button(bottom, text="  Simpan  ", command=self.save_changes, bg="#1a7f37", fg="white",
                  font=("Segoe UI", 9, "bold"), relief="flat", padx=8, pady=2,
                  activebackground="#15692d", activeforeground="white").pack(side="right", padx=(3, 3))

    def show_print(self, report):
        """Switch to the view page showing every line of this printout row (from db_prints)."""
        try:
            db_sync_one(report[6])  # the file may have been edited in Excel just now
            report = list(db_print_row(report[6]))
            items = db_report_items(report[6])
        except sqlite3.Error as e:
            messagebox.showerror(APP_TITLE, f"Gagal membaca riwayat:\n{e}", parent=self)
            return
        self._report = report
        self._pending = {}
        self._unlocked = not report[3]  # an approved printout needs a password before it can be changed
        self._update_view_hint()
        for key, value in zip(("waktu", "p_ecer", "p_kantor", "disetujui"),
                              (fmt_time(report[0]), *report[1:4])):
            self.view_vars[key].set(cell_text(value) or "-")
        self.vtree.delete(*self.vtree.get_children())
        for i, (line_no, *values) in enumerate(items):
            self.vtree.insert("", "end", iid=str(line_no), values=[cell_text(v) for v in values],
                              tags=(f"row{i % 2}",))
        self._show_view_info()
        self.list_page.pack_forget()
        self.view_page.pack(fill="both", expand=True)

    def _show_view_info(self):
        changed = sum(len(v) for v in self._pending.values())
        if changed:
            self.view_info.config(text=f"{changed} perubahan belum disimpan - klik Simpan.", fg="#a15c00")
        else:
            self.view_info.config(text=f"{len(self.vtree.get_children())} barang  -  "
                                       f"File: {os.path.basename(self._report[5] or '')}", fg="#333")

    def _confirm_pending(self):
        """Before leaving the printout: if there are unsaved changes ask Simpan / Jangan simpan / Batal.
        True = go on (saved or thrown away), False = stay."""
        if not self._pending:
            return True
        answer = messagebox.askyesnocancel(APP_TITLE, "Ada perubahan yang belum disimpan.\n\n"
                                                      "Simpan sekarang?\n(Yes = simpan, No = buang perubahan)",
                                           parent=self)
        if answer is None:
            return False
        if answer:
            return self.save_changes()
        self._pending = {}
        return True

    def back(self):
        """Back to the list of printouts, with the printout just viewed selected."""
        close_editors(self.vtree)
        if not self._confirm_pending():
            return
        report_id = self._report[6]
        self._report = None
        self.view_page.pack_forget()
        self.list_page.pack(fill="both", expand=True)
        self.refresh()  # values may have changed (totals per item)
        for iid in self.tree.get_children():
            if self.mode_var.get() == "prints" and self._rows[int(iid)][6] == report_id:
                self.tree.selection_set(iid)
                self.tree.see(iid)
                self.tree.focus(iid)
                break
        self.tree.focus_set()

    def _on_view_double_click(self, event):
        if self.vtree.identify_region(event.x, event.y) != "cell":
            return
        iid = self.vtree.identify_row(event.y)
        key = self.vtree["columns"][int(self.vtree.identify_column(event.x)[1:]) - 1]
        if iid and key in EDIT_KEYS and self._unlock():
            CellEditor(self.vtree, iid, key, EDIT_KEYS, self._save_cell)

    def _update_view_hint(self):
        text = ("Klik dua kali pada kolom J. Koli, Gudang, Stock Opname untuk mengisi atau mengubah "
                "lalu klik Simpan (menyimpan ke Riwayat dan ke file Excel).")
        if not self._unlocked:
            text = "Laporan ini sudah disetujui, untuk mengubah perlu password.  " + text
        self.view_hint.config(text=text, fg="#555" if self._unlocked else "#a15c00")

    def _unlock(self):
        """True if this printout may be changed. An approved one asks once for a password: the password of any
        person in Pengaturan Disetujui, or the settings password. It stays unlocked until Kembali."""
        if self._unlocked:
            return True
        pw = simpledialog.askstring(APP_TITLE, f"Laporan ini sudah disetujui oleh {self._report[3]}.\n"
                                               "Masukkan password untuk mengubah:", show="*", parent=self)
        if pw is None:
            return False
        pw = pw.strip()
        if not check_settings_password(self.app.cfg, pw) and not find_approver(self.app.cfg, pw):
            messagebox.showerror(APP_TITLE, "Password salah.", parent=self)
            return False
        self._unlocked = True
        self._update_view_hint()
        return True

    def _save_cell(self, iid, key, value):
        """A value typed in Riwayat: only shown and remembered; Simpan writes it (fast typing, one file write)."""
        self.vtree.set(iid, key, value)
        self._pending.setdefault(int(iid), {})[key] = value
        self._show_view_info()

    def save_changes(self):
        """Simpan: write the report's xlsx file first (so both stay the same), then Riwayat. True if saved."""
        close_editors(self.vtree)  # keep a value still being typed
        if not self._pending:
            self._show_view_info()
            return True
        keys = [c[0] for c in COLUMNS]
        rows = [dict(zip(keys, self.vtree.item(line, "values"))) for line in self.vtree.get_children()]
        self.config(cursor="watch")
        self.update_idletasks()
        try:
            mtime = self._write_file(self._report, rows, f"Diubah lewat Riwayat {datetime.now():%d/%m/%Y %H:%M}")
            if mtime is None:
                return False
            db_update_items(self._report[6], self._pending, mtime)
        except sqlite3.Error as e:
            messagebox.showerror(APP_TITLE, f"Gagal menyimpan:\n{e}", parent=self)
            return False
        finally:
            self.config(cursor="")
        self._pending = {}
        self._show_view_info()
        self.view_info.config(text="Tersimpan di Riwayat dan di file Excel.  " + self.view_info.cget("text"))
        return True

    def _write_file(self, report, rows, note):
        """(Re)write the xlsx of a printout row from rows (dicts of the 10 columns, as text).
        Returns the file's new 'last changed' time, or None (after a message) if it could not be written."""
        path = report[5]
        rows = [dict(r, stn=cell_value(r.get("stn", ""))) for r in rows]  # StnPerPak a number again
        printed = datetime.strptime(report[0], "%Y-%m-%d %H:%M:%S")
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            save_report(self.app.build_workbook(printed.strftime("%d/%m/%Y"), report[1] or "", report[2] or "",
                                                report[3] or "", rows, note=note), path)
        except PermissionError:
            messagebox.showwarning(APP_TITLE, "File laporan ini sedang dibuka di Excel, jadi tidak bisa diubah.\n"
                                              f"Tutup dulu file-nya di Excel, lalu coba lagi.\n\n{path}", parent=self)
            return None
        except Exception as e:  # noqa: BLE001
            messagebox.showerror(APP_TITLE, f"Gagal menyimpan file:\n{path}\n{e}", parent=self)
            return None
        return file_mtime(path)

    def print_selected(self):
        """Print button on the list: print the selected printout."""
        sel = self.tree.selection()
        if self.mode_var.get() != "prints" or not sel:
            messagebox.showinfo(APP_TITLE, "Klik dulu satu print di Daftar print.", parent=self)
            return
        self.print_report(self._rows[int(sel[0])])

    def print_report(self, report):
        """Print a printout's own xlsx file (it always matches Riwayat). A missing file is made again first."""
        close_editors(self.vtree)  # keep a value still being typed
        if self._report and self._pending:
            answer = messagebox.askyesnocancel(APP_TITLE, "Ada perubahan yang belum disimpan.\n\n"
                                                          "Simpan dulu lalu print?\n"
                                                          "(Yes = simpan lalu print, No = print tanpa perubahan)",
                                               parent=self)
            if answer is None or (answer and not self.save_changes()):
                return
        try:
            db_sync_one(report[6])  # pick up edits made in Excel before printing
            report = list(db_print_row(report[6]))
        except sqlite3.Error as e:
            messagebox.showerror(APP_TITLE, f"Gagal membaca riwayat:\n{e}", parent=self)
            return
        path = report[5]
        if not os.path.isfile(path):
            with closing(db_connect()) as con:
                rows = db_report_content(con, report[6])[3]
            mtime = self._write_file(report, rows, f"Dibuat ulang dari Riwayat {datetime.now():%d/%m/%Y %H:%M}")
            if mtime is None:
                return
            with closing(db_connect()) as con, con:
                con.execute("UPDATE reports SET file_mtime = ? WHERE id = ?", (mtime, report[6]))
        if self.app.send_out(path, parent=self):
            done = "dikirim ke printer" if self.app.auto_print_var.get() else "dibuka di Excel untuk di-print"
            info = self.view_info if self._report else self.info
            info.config(text=f"{os.path.basename(path)} {done}.", fg="#333")

    def open_file(self):
        path = self._report[5] if self._report else None
        if path and os.path.isfile(path):
            os.startfile(path)
        else:
            messagebox.showwarning(APP_TITLE, f"File laporan tidak ditemukan:\n{path}", parent=self)

    # ---------- history file ----------
    def start_new(self):
        if not messagebox.askyesno(
                APP_TITLE, "Mulai riwayat baru yang kosong?\n\n"
                           "Riwayat lama TIDAK dihapus, tetapi disimpan sebagai file cadangan "
                           "(laporan_backup_....db) di folder aplikasi.\n"
                           "File Excel laporan tidak diubah.", parent=self):
            return
        try:
            backup = db_start_new()
        except OSError as e:
            messagebox.showerror(APP_TITLE, f"Gagal membuat riwayat baru:\n{e}", parent=self)
            return
        messagebox.showinfo(APP_TITLE, f"Riwayat baru dimulai.\nRiwayat lama disimpan di:\n{backup}", parent=self)
        self.refresh()

    @staticmethod
    def _date(var):
        """Date typed in a Dari / Sampai box, or today if it can't be read."""
        try:
            return datetime.strptime(var.get().strip(), "%d/%m/%Y")
        except ValueError:
            return datetime.now()

    def import_old(self):
        pw = simpledialog.askstring(APP_TITLE, "Masukkan password pengaturan untuk impor laporan:", show="*",
                                    parent=self)
        if pw is None:
            return
        if not check_settings_password(self.app.cfg, pw):
            messagebox.showerror(APP_TITLE, "Password salah.", parent=self)
            return
        start = self.app.output_dir()
        paths = filedialog.askopenfilenames(
            title="Pilih file laporan (Ctrl+A = pilih semua)", parent=self,
            initialdir=start if os.path.isdir(start) else app_dir(),
            filetypes=[("Laporan Excel", "*.xlsx"), ("Semua file", "*.*")])
        if not paths:
            return
        self.config(cursor="watch")
        self.update_idletasks()
        try:
            added, updated, same, others, failed = import_report_files(paths)
        except sqlite3.Error as e:
            messagebox.showerror(APP_TITLE, f"Gagal impor:\n{e}", parent=self)
            return
        finally:
            self.config(cursor="")
        msg = f"{len(added)} laporan baru ditambahkan ke Riwayat."
        if updated:
            msg += f"\n{updated} laporan sudah ada dan diperbarui dari file Excel (file sudah diubah)."
        if same:
            msg += f"\n{same} laporan sudah ada dan sama dengan file (tidak berubah)."
        if others:
            msg += (f"\n\n{len(others)} file bukan laporan (nama file harus 'Laporan Stock Ecer "
                    "tanggal_jam.xlsx'):\n" + "\n".join(others[:10]))
        if failed:
            msg += f"\n\n{len(failed)} file gagal dibaca:\n" + "\n".join(failed[:10])
        if added:  # widen Dari / Sampai so the imported reports show in the list
            self.from_var.set(min(added + [self._date(self.from_var)]).strftime("%d/%m/%Y"))
            self.to_var.set(max(added + [self._date(self.to_var)]).strftime("%d/%m/%Y"))
        messagebox.showinfo(APP_TITLE, msg, parent=self)
        self.refresh()


def _fill(img, color, x0, y0, x1, y1):
    img.put(color, to=(x0, y0, x1, y1))


def folder_icon(master):
    """20x20 yellow folder, drawn here so no image file is needed next to the exe."""
    img = tk.PhotoImage(master=master, width=20, height=20)
    _fill(img, "#b8860b", 1, 3, 9, 7)     # tab outline
    _fill(img, "#e0a800", 2, 4, 8, 7)     # tab
    _fill(img, "#b8860b", 1, 6, 19, 18)   # body outline
    _fill(img, "#f4c430", 2, 7, 18, 17)   # body
    _fill(img, "#ffd966", 2, 9, 18, 17)   # front flap
    _fill(img, "#e0a800", 2, 9, 18, 10)   # flap edge
    return img


def save_icon(master):
    """20x20 blue floppy disk (save)."""
    img = tk.PhotoImage(master=master, width=20, height=20)
    _fill(img, "#173f80", 1, 1, 19, 19)   # outline
    _fill(img, "#2563c9", 2, 2, 18, 18)   # body
    _fill(img, "#d9e1ec", 5, 2, 14, 8)    # metal shutter
    _fill(img, "#2563c9", 10, 3, 12, 7)   # shutter slot
    _fill(img, "#ffffff", 4, 11, 16, 18)  # label
    _fill(img, "#9aa7b8", 6, 13, 14, 14)  # label lines
    _fill(img, "#9aa7b8", 6, 15, 14, 16)
    return img


def gear_icon(master):
    """20x20 gray gear (settings): 8 teeth around a ring with a hole in the middle."""
    img = tk.PhotoImage(master=master, width=20, height=20)
    for y in range(20):
        for x in range(20):
            dx, dy = x - 9.5, y - 9.5
            r = math.hypot(dx, dy)
            on_tooth = math.cos(8 * math.atan2(dy, dx)) > 0.3
            if 3 <= r <= (9.5 if on_tooth else 7):
                img.put("#5a6270", to=(x, y, x + 1, y + 1))
    return img


def refresh_icon(master):
    """16x16 round arrow (refresh): a ring open at the top right, with an arrow head at the gap."""
    img = tk.PhotoImage(master=master, width=16, height=16)
    color = "#2563c9"
    for y in range(16):
        for x in range(16):
            dx, dy = x - 7.5, y - 7.5
            angle = math.degrees(math.atan2(-dy, dx)) % 360  # 0 = right, 90 = up
            if 4.0 <= math.hypot(dx, dy) <= 6.6 and not 20 < angle < 75:
                img.put(color, to=(x, y, x + 1, y + 1))
    # arrow head pointing down at the end of the ring (right side, just under the gap)
    for row, (x0, x1) in enumerate([(9, 16), (10, 15), (11, 14), (12, 13)]):
        img.put(color, to=(x0, 4 + row, x1, 5 + row))
    return img


def print_button(parent, command):
    """Blue Print button, like PRINT in the main window."""
    return tk.Button(parent, text="  Print  ", command=command, bg="#1f6feb", fg="white",
                     font=("Segoe UI", 9, "bold"), relief="flat", padx=8, pady=2,
                     activebackground="#1a5fd0", activeforeground="white")


class Tooltip:
    """Small hint shown while the mouse is over a widget. text can be a function, so the hint stays current."""

    def __init__(self, widget, text):
        self.widget, self.text, self.tip = widget, text, None
        widget.bind("<Enter>", self.show, add="+")
        widget.bind("<Leave>", self.hide, add="+")
        widget.bind("<ButtonPress>", self.hide, add="+")

    def show(self, _e=None):
        if self.tip:
            return
        x = self.widget.winfo_rootx() + self.widget.winfo_width() // 2
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 4
        self.tip = tk.Toplevel(self.widget)
        self.tip.overrideredirect(True)
        self.tip.geometry(f"+{x}+{y}")
        text = self.text() if callable(self.text) else self.text
        tk.Label(self.tip, text=text, bg="#ffffe0", relief="solid", bd=1, padx=6, pady=2).pack()

    def hide(self, _e=None):
        if self.tip:
            self.tip.destroy()
            self.tip = None


class ApproverSettings(tk.Toplevel):
    """Gear window: the people who can tick 'Disetujui' (Nama + Password) and the Excel edit password.
    Passwords are stored only as hashes, so they can be set here but never shown again.
    Changes are kept here until Simpan (on_save(approvers, excel_password), excel_password None = unchanged,
    '' = no lock); Batal / window X throws them away."""

    def __init__(self, master, approvers, excel_locked, on_save):
        super().__init__(master)
        self.title("Pengaturan")
        self.transient(master)
        self.resizable(False, False)
        self.on_save = on_save
        self.rows = [dict(a) for a in approvers]  # work on a copy

        tk.Label(self, text="Orang yang boleh mencentang \"Disetujui\". Password dipakai saat mencentang.",
                 fg="#555").pack(anchor="w", padx=12, pady=(10, 4))
        frame = tk.Frame(self)
        frame.pack(fill="both", padx=12)
        self.tree = ttk.Treeview(frame, columns=("nama", "password"), show="headings", height=8,
                                 selectmode="browse")
        self.tree.heading("nama", text="Nama")
        self.tree.heading("password", text="Password")
        self.tree.column("nama", width=220)
        self.tree.column("password", width=160)
        ys = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=ys.set)
        self.tree.pack(side="left")
        ys.pack(side="left", fill="y")
        for i, color in enumerate(ROW_COLORS):
            self.tree.tag_configure(f"row{i}", background=color)
        self.tree.bind("<<TreeviewSelect>>", self._on_select)

        form = tk.Frame(self)
        form.pack(fill="x", padx=12, pady=(8, 0))
        tk.Label(form, text="Orang baru: klik Baru, isi Nama dan Password, lalu klik Tambah.\n"
                            "Mengubah / menghapus: klik dulu nama di tabel. Ubah dengan Password kosong =\n"
                            "password lama tetap dipakai. Password disimpan teracak, tidak bisa dilihat lagi.",
                 fg="#555", justify="left").grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 4))
        form.columnconfigure(3, weight=1)  # the help text's extra width goes right, not between label and box
        tk.Label(form, text="Nama").grid(row=1, column=0, sticky="w")
        self.nama_var = tk.StringVar()
        self.nama_entry = ttk.Entry(form, textvariable=self.nama_var, width=30)
        self.nama_entry.grid(row=1, column=1, sticky="w", padx=4, pady=2)
        tk.Label(form, text="Password").grid(row=2, column=0, sticky="w")
        self.pw_var = tk.StringVar()
        ttk.Entry(form, textvariable=self.pw_var, width=30).grid(row=2, column=1, sticky="w", padx=4, pady=2)
        edit = tk.Frame(form)
        edit.grid(row=1, column=2, rowspan=3, padx=(8, 0))
        ttk.Button(edit, text="Baru", command=self.new).pack(fill="x", pady=1)
        ttk.Button(edit, text="Tambah", command=self.add).pack(fill="x", pady=1)
        ttk.Button(edit, text="Ubah", command=self.change).pack(fill="x", pady=1)
        ttk.Button(edit, text="Hapus", command=self.delete).pack(fill="x", pady=1)

        excel = ttk.LabelFrame(self, text="Password Excel", padding=(8, 4))
        excel.pack(fill="x", padx=12, pady=(12, 0))
        tk.Label(excel, text="File laporan bisa dibuka dan di-print siapa saja, tetapi untuk mengedit perlu "
                            "password ini.\nIsi Password baru untuk mengganti, kosongkan = tidak diubah. "
                            "Catat password-nya, tidak bisa dilihat lagi.", fg="#555", justify="left"
                 ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 4))
        self.excel_locked = excel_locked
        self.excel_remove = False
        self.excel_status = tk.Label(excel, font=("Segoe UI", 9, "bold"))
        self.excel_status.grid(row=1, column=0, columnspan=3, sticky="w")
        tk.Label(excel, text="Password baru").grid(row=2, column=0, sticky="w")
        self.excel_pw_var = tk.StringVar()
        ttk.Entry(excel, textvariable=self.excel_pw_var, width=30).grid(row=2, column=1, sticky="w", padx=4, pady=2)
        ttk.Button(excel, text="Matikan kunci", command=self.remove_excel_lock).grid(row=2, column=2, sticky="w")
        excel.columnconfigure(2, weight=1)  # the wide help text's extra space goes right, not between label and box
        self._show_excel_status()

        sp = ttk.LabelFrame(self, text="Password Pengaturan", padding=(8, 4))
        sp.pack(fill="x", padx=12, pady=(12, 0))
        tk.Label(sp, text="Password untuk membuka jendela ini, Impor Laporan Lama, dan mengubah laporan yang sudah\n"
                          "disetujui. Kosongkan = tidak diubah. Catat password-nya, tidak bisa dilihat lagi.",
                 fg="#555", justify="left").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 4))
        self.settings_pw_var, self.settings_pw2_var = tk.StringVar(), tk.StringVar()
        for row, (label, var) in enumerate([("Password baru", self.settings_pw_var),
                                            ("Ulangi password", self.settings_pw2_var)], start=1):
            tk.Label(sp, text=label).grid(row=row, column=0, sticky="w")
            ttk.Entry(sp, textvariable=var, width=30, show="*").grid(row=row, column=1, sticky="w", padx=4, pady=2)
        sp.columnconfigure(2, weight=1)

        bottom = tk.Frame(self)
        bottom.pack(fill="x", padx=12, pady=10)
        ttk.Button(bottom, text="Batal", command=self.destroy).pack(side="right")
        ttk.Button(bottom, text="Simpan", command=self.save).pack(side="right", padx=4)
        self.bind("<Escape>", lambda e: self.destroy())

        self._render()
        self.grab_set()  # finish here before going back to the report
        self.focus_force()
        self.nama_entry.focus_set()

    def _render(self, select=None):
        self.tree.delete(*self.tree.get_children())
        for i, a in enumerate(self.rows):
            self.tree.insert("", "end", iid=str(i), values=(a["nama"], "••••••"), tags=(f"row{i % 2}",))
        if select is not None and 0 <= select < len(self.rows):
            self.tree.selection_set(str(select))
            self.tree.see(str(select))

    def _selected(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def _on_select(self, _e=None):
        i = self._selected()
        if i is not None:
            self.nama_var.set(self.rows[i]["nama"])
            self.pw_var.set("")  # the password can't be shown; empty = keep it

    def _form(self, skip=None):
        """{'nama', 'hash'} from the boxes, or None (with a message) if they can't be used.
        skip = the row being changed: its password may stay empty (keeps the old hash)."""
        nama, pw = self.nama_var.get().strip(), self.pw_var.get().strip()
        if not nama or (not pw and skip is None):
            messagebox.showwarning(APP_TITLE, "Nama dan Password harus diisi.", parent=self)
            return None
        if not pw:
            return {"nama": nama, "hash": self.rows[skip]["hash"]}
        if any(check_secret(pw, a["hash"]) for j, a in enumerate(self.rows) if j != skip):
            messagebox.showwarning(APP_TITLE, "Password ini sudah dipakai orang lain. Pakai password lain.",
                                   parent=self)
            return None
        return {"nama": nama, "hash": hash_secret(pw)}

    def new(self):
        """Empty boxes for a new person (nothing selected in the table)."""
        self.tree.selection_remove(*self.tree.selection())
        self.nama_var.set("")
        self.pw_var.set("")
        self.nama_entry.focus_set()

    def add(self):
        entry = self._form()
        if entry:
            self.rows.append(entry)
            self._render()
            self.new()  # ready for the next person

    def change(self):
        i = self._selected()
        if i is None:
            messagebox.showinfo(APP_TITLE, "Klik dulu nama di tabel yang mau diubah.", parent=self)
            return
        entry = self._form(skip=i)
        if entry:
            self.rows[i] = entry
            self._render(select=i)

    def delete(self):
        i = self._selected()
        if i is None:
            messagebox.showinfo(APP_TITLE, "Klik dulu nama di tabel yang mau dihapus.", parent=self)
            return
        if messagebox.askyesno(APP_TITLE, f"Hapus {self.rows[i]['nama']}?", parent=self):
            del self.rows[i]
            self.nama_var.set("")
            self.pw_var.set("")
            self._render()

    def _show_excel_status(self):
        if self.excel_remove:
            text, color = "Kunci akan dimatikan setelah Simpan.", "#a15c00"
        elif self.excel_locked:
            text, color = "Status: file Excel dikunci (password sudah diatur).", "#0a7a2f"
        else:
            text, color = "Status: file Excel tidak dikunci.", "#555"
        self.excel_status.config(text=text, fg=color)

    def remove_excel_lock(self):
        self.excel_pw_var.set("")
        self.excel_remove = True
        self._show_excel_status()

    def save(self):
        settings_pw, again = self.settings_pw_var.get().strip(), self.settings_pw2_var.get().strip()
        if settings_pw != again:
            messagebox.showwarning(APP_TITLE, "Password Pengaturan baru dan ulangannya tidak sama.\n"
                                              "Ketik lagi keduanya (atau kosongkan untuk tidak mengubah).", parent=self)
            return
        new_pw = self.excel_pw_var.get().strip()
        self.on_save(self.rows, new_pw or ("" if self.excel_remove else None), settings_pw or None)
        self.destroy()


def close_editors(tree):
    """Save and close any typing box open in this table (before it is refilled or scrolled)."""
    for w in tree.winfo_children():
        if isinstance(w, CellEditor):
            w._finish(None)


class CellEditor(ttk.Combobox):
    """Type straight into a table cell. Enter = save and go down, Tab / Shift+Tab = next / previous
    column, Esc = cancel, clicking elsewhere = save.
    Columns in CHOICES get a drop-down list instead (picking a value saves it); the others are a
    plain typing box (made with ttk.Entry's constructor; a Combobox is an Entry with a list)."""

    def __init__(self, tree, iid, key, keys, on_save):
        close_editors(tree)  # only one box at a time
        self.tree, self.iid, self.key, self.keys, self.on_save = tree, iid, key, keys, on_save
        self.var = tk.StringVar(value=tree.set(iid, key))
        self.choices = CHOICES.get(key)
        if self.choices:
            ttk.Combobox.__init__(self, tree, textvariable=self.var, values=self.choices,
                                  state="readonly", justify="center")
            self.bind("<<ComboboxSelected>>", lambda e: self._finish(None))
            self.bind("<Delete>", lambda e: self.var.set(""))
            self.bind("<BackSpace>", lambda e: self.var.set(""))
        else:
            ttk.Entry.__init__(self, tree, textvariable=self.var, justify="center")
        self._done = False
        tree.see(iid)
        tree.update_idletasks()
        box = tree.bbox(iid, key)
        if not box:
            self.destroy()
            return
        x, y, w, h = box
        self.place(x=x, y=y, width=w, height=h)
        # the table's own click handling runs after a double-click and takes the keyboard focus
        # back, so give focus to the box only once that is done
        self.after_idle(self._take_focus)
        self.bind("<Return>", lambda e: self._finish("down"))
        self.bind("<KP_Enter>", lambda e: self._finish("down"))
        self.bind("<Tab>", lambda e: self._finish("right"))
        self.bind("<Shift-Tab>", lambda e: self._finish("left"))
        self.bind("<Escape>", lambda e: self._cancel())
        self.bind("<FocusOut>", lambda e: self.after(10, self._check_focus))

    def _check_focus(self):
        """Save when the focus left the box, but not when it only moved into the box's own drop-down list."""
        if self._done or not self.winfo_exists():
            return
        focus = str(self.tk.call("focus"))
        if focus == str(self) or focus.startswith(str(self) + "."):
            return
        self._finish(None)

    def _take_focus(self):
        if self._done or not self.winfo_exists():
            return
        self.focus_force()
        self.select_range(0, "end")
        self.icursor("end")
        if self.choices:
            self.tk.call("ttk::combobox::Post", self)  # open the list right away

    def _cancel(self):
        self._done = True
        self.destroy()
        self.tree.focus_set()
        return "break"

    def _finish(self, move):
        if self._done:
            return "break"
        self._done = True
        tree, iid, key, keys = self.tree, self.iid, self.key, self.keys
        value = self.var.get().strip()
        self.destroy()
        if not tree.exists(iid):
            return "break"
        if value != tree.set(iid, key):
            self.on_save(iid, key, value)
        nxt_iid, nxt_key = None, key
        if move == "down":
            nxt_iid = tree.next(iid)
        elif move in ("right", "left"):
            i = keys.index(key) + (1 if move == "right" else -1)
            if 0 <= i < len(keys):
                nxt_iid, nxt_key = iid, keys[i]
            elif i >= len(keys):  # past the last column: first column of the next row
                nxt_iid, nxt_key = tree.next(iid), keys[0]
            else:
                nxt_iid, nxt_key = tree.prev(iid), keys[-1]
        if nxt_iid:
            tree.selection_set(nxt_iid)
            CellEditor(tree, nxt_iid, nxt_key, keys, self.on_save)
        elif move:
            tree.focus_set()
        return "break"


class ItemPicker(tk.Toplevel):
    """Drop-down with a search bar on top and a list of all items."""

    def __init__(self, master, items, first_key, x, y, on_pick):
        super().__init__(master)
        self.withdraw()  # stay hidden until placed below the clicked header
        self.items = items
        self.on_pick = on_pick
        self.overrideredirect(True)
        self.transient(master)
        self.configure(bg="#888", padx=1, pady=1)

        # show the column that was clicked first
        order = [first_key] + [k for k in ("kode", "nama", "satuan", "stn") if k != first_key]
        self.order = order
        titles = {"kode": "Kode Barang", "nama": "Nama Barang", "satuan": "Satuan", "stn": "StnPerPak"}
        widths = {"kode": 110, "nama": 260, "satuan": 70, "stn": 80}

        body = tk.Frame(self, bg="white", padx=6, pady=6)
        body.pack(fill="both", expand=True)
        tk.Label(body, text="Cari:", bg="white").grid(row=0, column=0, sticky="w")
        self.search_var = tk.StringVar()
        self.search = ttk.Entry(body, textvariable=self.search_var, width=50)
        self.search.grid(row=0, column=1, sticky="ew", padx=(4, 0))

        self.tree = ttk.Treeview(body, columns=order, show="headings", height=15, selectmode="browse")
        for k in order:
            self.tree.heading(k, text=titles[k])
            self.tree.column(k, width=widths[k], anchor="w" if k in ("kode", "nama") else "center")
        sb = ttk.Scrollbar(body, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(6, 0))
        sb.grid(row=1, column=2, sticky="ns", pady=(6, 0))
        self.info = tk.Label(body, text="", bg="white", fg="#666", anchor="w")
        self.info.grid(row=2, column=0, columnspan=3, sticky="ew")
        body.columnconfigure(1, weight=1)
        body.rowconfigure(1, weight=1)

        self._after = None
        self._shown = []
        self.search_var.trace_add("write", lambda *a: self._schedule_filter())
        self.search.bind("<Down>", self._focus_list)
        self.search.bind("<Return>", lambda e: self._pick_current())
        self.tree.bind("<Double-1>", lambda e: self._pick_current())
        self.tree.bind("<Return>", lambda e: self._pick_current())
        self.bind("<Escape>", lambda e: (self.close(), self.master.focus_force()))
        self.bind("<FocusOut>", self._on_focus_out)

        self._filter()
        self.update_idletasks()
        # keep the drop-down on screen
        w, h = self.winfo_reqwidth(), self.winfo_reqheight()
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        x = max(0, min(x, sw - w))
        y = max(0, min(y, sh - h - 40))
        # on Windows, lift() on a borderless window that is not shown yet moves it to 0,0,
        # so place and show it first, then bring it to the front
        self.geometry(f"+{x}+{y}")
        self.deiconify()
        self.update_idletasks()
        self.lift()
        self.attributes("-topmost", True)
        self.search.focus_force()

    def _schedule_filter(self):
        if self._after:
            self.after_cancel(self._after)
        self._after = self.after(150, self._filter)

    def _filter(self):
        self._after = None
        terms = self.search_var.get().lower().split()
        matches = [it for it in self.items if all(t in it["_search"] for t in terms)]
        self.tree.delete(*self.tree.get_children())
        self._shown = matches[:MAX_DROPDOWN_ROWS]
        for i, it in enumerate(self._shown):
            self.tree.insert("", "end", iid=str(i), values=[it[k] for k in self.order])
        if self._shown:
            self.tree.selection_set("0")
            self.tree.focus("0")
        total = len(matches)
        if total > MAX_DROPDOWN_ROWS:
            self.info.config(text=f"{total} barang cocok (menampilkan {MAX_DROPDOWN_ROWS} pertama, ketik untuk mempersempit)")
        else:
            self.info.config(text=f"{total} barang cocok")

    def _focus_list(self, _e):
        if self._shown:
            self.tree.focus_set()
            sel = self.tree.selection() or ("0",)
            self.tree.focus(sel[0])
        return "break"

    def _pick_current(self):
        sel = self.tree.selection()
        if not sel:
            return
        item = self._shown[int(sel[0])]
        self.close()
        # Windows does not hand the keyboard back to the main window when this borderless window closes,
        # so the next click there (e.g. in Perwakilan) would not get the typing; take it back here
        self.master.focus_force()
        self.on_pick(item)

    def _on_focus_out(self, _e):
        # close when focus leaves the drop-down entirely
        self.after(100, self._check_focus)

    def _check_focus(self):
        try:
            f = self.focus_get()
        except (KeyError, tk.TclError):
            f = None
        if f is None or not str(f).startswith(str(self)):
            self.close()

    def close(self):
        if self.winfo_exists():
            self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1280x640")
        self.minsize(900, 450)
        self.cfg = load_config()
        upgrade_config(self.cfg)
        self.items = []
        self.rows = []  # list of item dicts in the report
        self.picker = None
        self.history = None

        style = ttk.Style(self)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        style.configure("Report.Treeview", rowheight=24)
        style.configure("Report.Treeview.Heading", font=("Segoe UI", 9, "bold"))

        self._build_ui()
        self.after(50, self._initial_load)

    # ---------- UI ----------
    def _build_ui(self):
        tk.Label(self, text=APP_TITLE, font=("Segoe UI", 16, "bold")).pack(pady=(10, 4))

        top = tk.Frame(self)
        top.pack(fill="x", padx=14)
        info = tk.Frame(top)
        info.pack(side="left")
        tk.Label(info, text="Tanggal", width=18, anchor="w").grid(row=0, column=0, sticky="w")
        self.tanggal_var = tk.StringVar(value=datetime.now().strftime("%d/%m/%Y"))
        tk.Label(info, textvariable=self.tanggal_var, anchor="w").grid(row=0, column=1, sticky="w")
        tk.Label(info, text="Perwakilan Ecer", width=18, anchor="w").grid(row=1, column=0, sticky="w")
        self.ecer_var = tk.StringVar()
        e = ttk.Entry(info, textvariable=self.ecer_var, width=35)
        e.grid(row=1, column=1, sticky="w", pady=2)
        e.bind("<Button-1>", lambda ev: ev.widget.focus_force(), add="+")  # a click always gets the typing
        tk.Label(info, text="Perwakilan Kantor", width=18, anchor="w").grid(row=2, column=0, sticky="w")
        self.kantor_var = tk.StringVar()
        e = ttk.Entry(info, textvariable=self.kantor_var, width=35)
        e.grid(row=2, column=1, sticky="w", pady=2)
        e.bind("<Button-1>", lambda ev: ev.widget.focus_force(), add="+")

        # settings as icons only; hover shows the current file / folder
        icons = tk.Frame(top)
        icons.pack(side="right", anchor="n")
        self._icons = {"folder": folder_icon(self), "save": save_icon(self),
                       "gear": gear_icon(self)}  # keep references for Tk
        b = ttk.Button(icons, image=self._icons["folder"], command=self.choose_db)
        b.pack(side="left", padx=2)
        Tooltip(b, lambda: f"Pilih file database\nSekarang: {self.cfg.get('db_path') or '(belum dipilih)'}")
        b = ttk.Button(icons, image=self._icons["save"], command=self.choose_output_dir)
        b.pack(side="left", padx=2)
        Tooltip(b, lambda: f"Pilih folder untuk menyimpan laporan\nSekarang: {self.output_dir()}")
        b = ttk.Button(icons, image=self._icons["gear"], command=self.open_settings)
        b.pack(side="left", padx=2)
        Tooltip(b, "Pengaturan: Disetujui (nama dan password) dan password Excel")

        btns = tk.Frame(self)
        btns.pack(fill="x", padx=14, pady=(8, 0))
        ttk.Button(btns, text="Hapus Baris", command=self.delete_selected).pack(side="left", padx=3)
        ttk.Button(btns, text="Kosongkan", command=self.clear_rows).pack(side="left", padx=3)
        ttk.Button(btns, text="Riwayat", command=self.open_history).pack(side="left", padx=(20, 3))
        tk.Button(btns, text="  PRINT  ", command=self.print_report, bg="#1f6feb", fg="white",
                  font=("Segoe UI", 10, "bold"), relief="flat", padx=10, pady=4,
                  activebackground="#1a5fd0", activeforeground="white").pack(side="right")
        # small "otomatis" tick next to PRINT: ticked = straight to the printer, else open the file in Excel
        auto = tk.Frame(btns)
        auto.pack(side="right", padx=(0, 6))
        self.auto_print_var = tk.BooleanVar(value=bool(self.cfg.get("auto_print")))
        b = ttk.Checkbutton(auto, variable=self.auto_print_var, command=self._save_auto_print)
        b.pack()
        tk.Label(auto, text="otomatis", font=("Segoe UI", 7), fg="#555").pack()
        Tooltip(b, "Dicentang: PRINT langsung ke printer.\nTidak dicentang: file dibuka di Excel, print sendiri.")
        self.approved_by = ""
        self.approve_var = tk.BooleanVar(value=False)
        self.approve_name_var = tk.StringVar()
        tk.Label(btns, textvariable=self.approve_name_var, fg="#0a7a2f",
                 font=("Segoe UI", 10, "bold")).pack(side="right", padx=(0, 16))
        ttk.Checkbutton(btns, text="Disetujui", variable=self.approve_var,
                        command=self.toggle_approval).pack(side="right", padx=(0, 4))

        tk.Label(self, text="Klik tanda  ＋  pada judul kolom untuk menambah barang.  "
                            "Klik dua kali pada kolom J. Koli, Gudang, Stock Opname untuk mengisi (Enter / Tab untuk pindah).  "
                            "Klik dua kali pada baris untuk mengganti barang.  Tombol Delete untuk menghapus baris.",
                 fg="#555").pack(anchor="w", padx=14, pady=(8, 2))

        frame = tk.Frame(self)
        frame.pack(fill="both", expand=True, padx=14, pady=(0, 6))
        keys = [c[0] for c in COLUMNS]
        self.tree = ttk.Treeview(frame, columns=keys, show="headings", style="Report.Treeview",
                                 selectmode="extended")
        for key, text, width, _xw, src in COLUMNS:
            label = text.replace("\n", " ")
            if src:
                label += "  ＋"
                self.tree.heading(key, text=label, command=lambda k=key: self.open_picker(k))
            else:
                self.tree.heading(key, text=label)
            self.tree.column(key, width=width, minwidth=50,
                             anchor="w" if key in ("kode", "nama") else "center")
        ys = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        xs = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ys.grid(row=0, column=1, sticky="ns")
        xs.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", self._on_double_click)
        self.tree.bind("<MouseWheel>", lambda e: close_editors(self.tree), add="+")
        self.tree.bind("<Delete>", lambda e: self.delete_selected())
        for i, color in enumerate(ROW_COLORS):
            self.tree.tag_configure(f"row{i}", background=color)

        self.status_var = tk.StringVar()
        tk.Label(self, textvariable=self.status_var, anchor="w", fg="#333", relief="sunken", bd=1
                 ).pack(fill="x", side="bottom")

    # ---------- database ----------
    def _initial_load(self):
        get_approvers(self.cfg)  # first run: bring the names from the old password_disetujui.txt
        # first run with history: bring in the reports that were printed before the database existed
        if not os.path.isfile(DB_PATH):
            try:
                added, _failed = import_report_folder(self.output_dir())
            except sqlite3.Error:
                added = 0
            if added:
                messagebox.showinfo(APP_TITLE, f"{added} laporan lama dimasukkan ke Riwayat.")
        candidates = [self.cfg.get("db_path"), os.path.join(app_dir(), DEFAULT_DB_NAME)]
        for path in candidates:
            if path and os.path.isfile(path):
                self.load_db(path)
                return
        messagebox.showinfo(APP_TITLE, f"File database \"{DEFAULT_DB_NAME}\" tidak ditemukan.\n"
                                       "Silakan pilih file database (xlsx).")
        self.choose_db()

    def output_dir(self):
        return self.cfg.get("output_dir") or os.path.join(app_dir(), "Laporan")

    def choose_output_dir(self):
        path = filedialog.askdirectory(title="Pilih folder untuk menyimpan laporan",
                                       initialdir=self.output_dir() if os.path.isdir(self.output_dir()) else app_dir())
        if path:
            self.cfg["output_dir"] = os.path.normpath(path)
            save_config(self.cfg)
            self.status_var.set(f"Laporan akan disimpan di: {self.cfg['output_dir']}")

    def open_settings(self):
        close_editors(self.tree)  # save a value still being typed first
        pw = simpledialog.askstring(APP_TITLE, "Masukkan password pengaturan:", show="*", parent=self)
        if pw is None:
            return
        if not check_settings_password(self.cfg, pw):
            messagebox.showerror(APP_TITLE, "Password salah.")
            return
        ApproverSettings(self, get_approvers(self.cfg), bool(self.cfg.get("excel_password_hash")),
                         self._save_settings)

    def _save_settings(self, approvers, excel_password, settings_password=None):
        """excel_password: None = keep the current lock, '' = no lock, else the new password.
        settings_password: None = unchanged, else the new settings password (only its hash is kept)."""
        self.cfg["approvers"] = approvers
        if excel_password == "":
            self.cfg.pop("excel_password_hash", None)
        elif excel_password:
            self.cfg["excel_password_hash"] = excel_hash(excel_password)
        if settings_password:
            self.cfg["settings_password_hash"] = hash_secret(settings_password)
        save_config(self.cfg)
        lock = "file Excel dikunci" if self.cfg.get("excel_password_hash") else "file Excel tidak dikunci"
        changed = ", password pengaturan diganti" if settings_password else ""
        self.status_var.set(f"Pengaturan disimpan: {len(approvers)} orang Disetujui, {lock}{changed}.")

    def choose_db(self):
        current = self.cfg.get("db_path")
        start = os.path.dirname(current) if current and os.path.isdir(os.path.dirname(current)) else app_dir()
        path = filedialog.askopenfilename(title="Pilih file database", initialdir=start,
                                          filetypes=[("Excel", "*.xlsx *.xlsm"), ("Semua file", "*.*")])
        if path:
            self.load_db(path)

    def load_db(self, path):
        self.status_var.set("Memuat database...")
        self.config(cursor="watch")
        self.update_idletasks()
        try:
            self.items = load_items(path)
        except Exception as e:  # noqa: BLE001 - show any read error to the user
            messagebox.showerror(APP_TITLE, f"Gagal membaca database:\n{e}")
            self.status_var.set("Database belum dimuat.")
            return
        finally:
            self.config(cursor="")
        self.cfg["db_path"] = path
        save_config(self.cfg)
        self.status_var.set(f"Database: {os.path.basename(path)}  -  {len(self.items)} barang")

    # ---------- rows ----------
    def _column_x(self, key):
        x = 0
        for k in self.tree["columns"]:
            if k == key:
                return x
            x += self.tree.column(k, "width")
        return 0

    def open_picker(self, key, replace_index=None):
        close_editors(self.tree)  # save a value still being typed first
        if not self.items:
            messagebox.showwarning(APP_TITLE, "Database belum dimuat.")
            return
        if self.picker is not None and self.picker.winfo_exists():
            self.picker.close()
        x = self.tree.winfo_rootx() + self._column_x(key) - int(self.tree.xview()[0] * self._total_width())
        y = self.tree.winfo_rooty() + 24
        if replace_index is not None:
            bbox = self.tree.bbox(self.tree.get_children()[replace_index])
            if bbox:
                y = self.tree.winfo_rooty() + bbox[1] + bbox[3]
        self.picker = ItemPicker(self, self.items, key, x, y,
                                 lambda item: self.add_item(item, replace_index))

    def _total_width(self):
        return sum(self.tree.column(k, "width") for k in self.tree["columns"])

    def add_item(self, item, replace_index=None):
        for i, r in enumerate(self.rows):
            if r["kode"] == item["kode"] and i != replace_index:
                self.status_var.set(f"{item['kode']} sudah ada di laporan.")
                iid = self.tree.get_children()[i]
                self.tree.selection_set(iid)
                self.tree.see(iid)
                return
        # copy only the database fields, so typed values never change the loaded item list
        row = {k: item[k] for k in DB_KEYS}
        if replace_index is None:
            row.update({k: "" for k in HAND_KEYS})
            self.rows.append(row)
        else:  # replacing the item keeps what was already typed in that row
            row.update({k: self.rows[replace_index].get(k, "") for k in HAND_KEYS})
            self.rows[replace_index] = row
        self.refresh()
        idx = len(self.rows) - 1 if replace_index is None else replace_index
        iid = self.tree.get_children()[idx]
        self.tree.selection_set(iid)
        self.tree.see(iid)
        self.status_var.set(f"{item['kode']} - {item['nama']} ditambahkan.  Total {len(self.rows)} barang.")

    def refresh(self):
        close_editors(self.tree)
        self.tree.delete(*self.tree.get_children())
        for i, r in enumerate(self.rows):
            self.tree.insert("", "end", values=[r.get(c[0], "") for c in COLUMNS], tags=(f"row{i % 2}",))

    def _on_double_click(self, event):
        if self.tree.identify_region(event.x, event.y) != "cell":
            return
        iid = self.tree.identify_row(event.y)
        if not iid:
            return
        col = self.tree.identify_column(event.x)
        key = self.tree["columns"][int(col[1:]) - 1]
        if key in EDIT_KEYS:
            CellEditor(self.tree, iid, key, EDIT_KEYS, self._save_cell)
        elif key in DB_KEYS:
            self.open_picker(key, self.tree.index(iid))

    def _save_cell(self, iid, key, value):
        self.rows[self.tree.index(iid)][key] = value
        self.tree.set(iid, key, value)

    def delete_selected(self):
        close_editors(self.tree)  # save a value still being typed first
        sel = self.tree.selection()
        if not sel:
            return
        idxs = sorted((self.tree.index(i) for i in sel), reverse=True)
        for i in idxs:
            del self.rows[i]
        self.refresh()
        self.status_var.set(f"{len(idxs)} baris dihapus.  Total {len(self.rows)} barang.")

    def clear_rows(self):
        close_editors(self.tree)  # save a value still being typed first
        if self.rows and messagebox.askyesno(APP_TITLE, "Hapus semua barang dari laporan?"):
            self.rows.clear()
            self.refresh()

    # ---------- output ----------
    def build_workbook(self, tanggal, ecer, kantor, approved_by, rows, note=""):
        """The report as an xlsx workbook. note (e.g. 'Diubah lewat Riwayat ...') goes in the empty row 2."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Stock Ecer"
        ncol = len(COLUMNS)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        font = Font(name="Arial", size=11)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.cell(1, 1, "Laporan Stock Ecer").font = Font(name="Arial", size=12, bold=True)
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        if note:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
            ws.cell(2, 1, note).font = Font(name="Arial", size=9, italic=True)
            ws.cell(2, 1).alignment = Alignment(horizontal="center")

        for r, (label, value) in enumerate([("Tanggal", tanggal),
                                            ("Perwakilan Ecer", ecer),
                                            ("Perwakilan Kantor", kantor)]
                                           + ([("Disetujui oleh", approved_by)] if approved_by else []),
                                           start=3):
            ws.cell(r, 1, label).font = font
            c = ws.cell(r, 2, f": {value}")
            c.font = font

        # column width = widest text in that column (header lines, data, and the labels in column A);
        # empty hand-filled columns keep their default width as a minimum so there is room to write
        widths = []
        for ci, (key, text, _w, xw, src) in enumerate(COLUMNS, start=1):
            texts = text.split("\n") + [str(r.get(key, "")) for r in rows]
            if ci == 1:
                texts += ["Tanggal", "Perwakilan Ecer", "Perwakilan Kantor", "Disetujui oleh"]
            w = max(text_width(t) for t in texts) + 3
            widths.append(min(max(w, xw if not src else 8), MAX_COL_WIDTH))
            ws.column_dimensions[get_column_letter(ci)].width = widths[-1]

        head_row = 7
        for ci, (_k, text, _w, _xw, _s) in enumerate(COLUMNS, start=1):
            c = ws.cell(head_row, ci, text)
            c.font, c.alignment, c.border = font, center, border
        ws.row_dimensions[head_row].height = row_height(
            [c[1] for c in COLUMNS], widths, minimum=32)

        n = max(MIN_PRINT_ROWS, len(rows))
        # every ROWS_PER_PAGE item rows are followed by one hidden empty row, so each group can be its own
        # print area (see the page setup below); readers skip it because its Kode Barang is empty
        groups = []  # (first row, last row) of each group of item rows
        for i in range(n):
            r = head_row + 1 + i + i // ROWS_PER_PAGE
            if i % ROWS_PER_PAGE == 0:
                if i:
                    ws.row_dimensions[r - 1].hidden = True
                groups.append([r, r])
            groups[-1][1] = r
            item = rows[i] if i < len(rows) else None
            values = []
            for ci, col in enumerate(COLUMNS, start=1):
                v = item.get(col[0], "") if item else None
                if v is not None and not col[4]:  # typed values: numbers as numbers, blank as empty cell
                    v = cell_value(v) if v != "" else None
                values.append("" if v is None else str(v))
                c = ws.cell(r, ci, v)
                c.font, c.border = font, border
                c.alignment = Alignment(horizontal="left" if col[0] in ("kode", "nama") else "center",
                                        vertical="center", wrap_text=True)
            # taller row only when a text is longer than the maximum column width and wraps
            ws.row_dimensions[r].height = row_height(values, widths, minimum=20)

        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        # at most ROWS_PER_PAGE item rows per paper: each group is a separate print area, and Excel always
        # starts a print area on a new paper (page breaks would be ignored by "fit to page width", which is
        # needed because printers draw the columns wider or narrower). Areas that touch are merged by Excel,
        # hence the hidden row between groups. The top part (rows 1-7) is repeated on every paper.
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = len(groups)  # each group fits one paper, also with tall (wrapped) rows
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.75  # room under the table for "Halaman 1 dari 2" (footer at 0.3")
        ws.page_margins.footer = 0.3
        ws.page_margins.header = HEADER_MARGIN_IN  # the watermark picture starts near the top edge of the paper
        ws.oddHeader.center.text = "&G"  # the watermark picture, added by add_watermark() after saving
        ws.HeaderFooter.scaleWithDoc = False  # the watermark always fills the page, whatever the table's scale
        # table size on screen in pixels (Calibri 11 default: 7 px per width unit; 96 px per 72 pt) so
        # save_report can make the watermark exactly as big as the table
        wb.table_px = (sum(round(w * 7) for w in widths),
                       round(sum((ws.row_dimensions[r].height or 15) for r in range(1, groups[-1][1] + 1)
                                 if not ws.row_dimensions[r].hidden) * 4 / 3))
        # a long report prints on more papers: every paper repeats the whole top part (title, Tanggal,
        # Perwakilan, Disetujui, column headers) and says which page it is
        ws.print_title_rows = f"1:{head_row}"
        ws.oddFooter.center.text = "Halaman &P dari &N"
        ws.oddFooter.center.size = 9
        last_col = get_column_letter(ncol)
        ws.print_area = [f"A{1 if k == 0 else first}:{last_col}{last}" for k, (first, last) in enumerate(groups)]

        # edit lock (set in the gear window): the report can be opened, viewed and printed, but editing needs
        # the password (Review > Unprotect Sheet); structure lock stops adding, deleting or renaming sheets
        # (the config keeps only the hash Excel stores in the file, never the password itself)
        pw_hash = self.cfg.get("excel_password_hash")
        if pw_hash:
            ws.protection.set_password(pw_hash, already_hashed=True)
            ws.protection.sheet = True
            wb.security = WorkbookProtection(lockStructure=True)
            wb.security.set_workbook_password(pw_hash, already_hashed=True)
        return wb

    def print_report(self):
        close_editors(self.tree)  # save a value still being typed first
        if not self.rows and not messagebox.askyesno(APP_TITLE, "Laporan masih kosong. Tetap print?"):
            return
        now = datetime.now()
        tanggal = now.strftime("%d/%m/%Y")
        self.tanggal_var.set(tanggal)

        out_dir = self.output_dir()
        try:
            os.makedirs(out_dir, exist_ok=True)
        except OSError as e:
            messagebox.showerror(APP_TITLE, f"Folder simpan tidak bisa dibuat:\n{out_dir}\n{e}\n\n"
                                            "Silakan pilih folder lain.")
            return
        path = os.path.join(out_dir, f"Laporan Stock Ecer {now:%Y%m%d_%H%M%S}.xlsx")
        try:
            save_report(self.build_workbook(tanggal, self.ecer_var.get().strip(), self.kantor_var.get().strip(),
                                            self.approved_by, self.rows), path)
        except Exception as e:  # noqa: BLE001
            messagebox.showerror(APP_TITLE, f"Gagal menyimpan file:\n{e}")
            return
        self.status_var.set(f"Tersimpan: {path}")
        auto = self.auto_print_var.get()
        if not self.send_out(path):
            return
        try:
            db_save_report(now, self.ecer_var.get().strip(), self.kantor_var.get().strip(), path, self.rows,
                           self.approved_by)
        except sqlite3.Error as e:
            messagebox.showwarning(APP_TITLE, f"Laporan sudah dikirim ke printer, tetapi gagal dicatat di Riwayat:\n{e}")
            return
        finally:
            self.set_approval("")  # each report needs its own approval
        if auto:
            messagebox.showinfo(APP_TITLE, f"File tersimpan dan dikirim ke printer default:\n{path}")
        else:  # Excel opening is the feedback; no extra pop-up
            self.status_var.set(f"Tersimpan dan dibuka di Excel: {path}")

    def send_out(self, path, parent=None):
        """Otomatis ticked: send the saved file straight to the default printer; otherwise open it in Excel
        so the user prints it. Returns False (after a message) if that failed."""
        auto = self.auto_print_var.get()
        try:
            os.startfile(path, "print") if auto else os.startfile(path)
        except OSError as e:
            what = "mengirim ke printer" if auto else "membuka file di Excel"
            messagebox.showwarning(APP_TITLE, f"File tersimpan di:\n{path}\n\n"
                                              f"Tetapi gagal {what}:\n{e}\n\n"
                                              "Pastikan Microsoft Excel terpasang dan printer default sudah diatur.",
                                   parent=parent or self)
            return False
        return True

    def _save_auto_print(self):
        self.cfg["auto_print"] = self.auto_print_var.get()
        save_config(self.cfg)

    # ---------- approval ----------
    def toggle_approval(self):
        close_editors(self.tree)  # save a value still being typed first
        if not self.approve_var.get():  # unticked
            self.set_approval("")
            return
        self.approve_var.set(False)  # stays unticked until a valid password is given
        approvers = get_approvers(self.cfg)
        if not approvers:
            messagebox.showwarning(APP_TITLE, "Belum ada orang yang diatur untuk menyetujui.\n\n"
                                              "Klik tombol gear (pengaturan) di kanan atas untuk menambah.")
            return
        pw = simpledialog.askstring(APP_TITLE, "Masukkan password untuk menyetujui laporan:", show="*", parent=self)
        if pw is None:
            return
        name = find_approver(self.cfg, pw.strip())
        if not name:
            messagebox.showerror(APP_TITLE, "Password salah.")
            return
        self.set_approval(name)

    def set_approval(self, name):
        self.approved_by = name
        self.approve_var.set(bool(name))
        self.approve_name_var.set(name)
        if name:
            self.status_var.set(f"Laporan disetujui oleh {name}.")

    def open_history(self):
        close_editors(self.tree)  # save a value still being typed first
        if self.history is not None and self.history.winfo_exists():
            self.history.lift()
            self.history.focus_force()
            return
        self.history = HistoryWindow(self)


if __name__ == "__main__":
    App().mainloop()
